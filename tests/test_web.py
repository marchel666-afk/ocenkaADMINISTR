from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app import db
from app.transcription import Segment, Transcript
from app.web import create_app, guess_datetime

from .conftest import FakeLLM, FakeTranscriber, sample_text


def make(settings, transcriber=None, llm=None):
    transcriber = transcriber or FakeTranscriber()
    llm = llm or FakeLLM()
    app = create_app(settings, transcriber=transcriber, llm=llm)
    return app, TestClient(app), transcriber, llm


def process_all(app):
    with app.state.session_factory() as s:
        ids = s.scalars(select(db.Call.id).where(db.Call.status == db.QUEUED).order_by(db.Call.id)).all()
    for call_id in ids:
        app.state.processor.process(call_id)


def get_call(app, call_id=1) -> db.Call:
    with app.state.session_factory() as s:
        return s.get(db.Call, call_id)


def test_upload_audio_full_flow(settings):
    app, client, transcriber, llm = make(settings)
    client.post("/employees", data={"name": "Анна Смирнова"})

    r = client.post(
        "/calls/upload",
        files=[("files", ("2026-10-08_14-32-05.mp3", b"fake-audio", "audio/mpeg"))],
        data={"direction": "in"},
    )
    assert r.status_code == 200  # после редиректа — список звонков
    call = get_call(app)
    assert call.status == db.QUEUED and call.audio_path
    assert call.started_at.isoformat() == "2026-10-08T14:32:05"
    assert (settings.audio_dir / call.audio_path).read_bytes() == b"fake-audio"

    process_all(app)
    call = get_call(app)
    assert call.status == db.DONE, call.status_message
    assert call.call_type == "incoming"
    assert call.employee_id is not None  # найден по имени «Анна» из разговора
    assert 0 < call.score_total < 100
    assert call.score_general is not None and call.score_scenario is not None
    assert call.llm_tokens == 3000 and call.llm_cost_usd == pytest.approx(0.002)
    assert transcriber.calls == 1
    with app.state.session_factory() as s:
        results = s.scalars(select(db.CriterionResult).where(db.CriterionResult.call_id == 1)).all()
    assert {r.criterion_id for r in results if r.result == "no"} == {"IN11", "G04"}

    page = client.get("/calls/1").text
    assert "Обратная связь администратору" in page
    assert "Нет регалий врача" in page
    assert "не оценивается по записи" in page
    assert "Анна Смирнова" in page

    assert "Анна Смирнова" in client.get("/calls").text
    stats = client.get("/stats").text
    assert "Анна Смирнова" in stats and "ФИО и регалии врача" in stats

    xlsx = client.get("/export.xlsx")
    wb = load_workbook(BytesIO(xlsx.content))
    assert wb.sheetnames == ["Звонки", "Пункты", "Сводка"]
    assert wb["Звонки"].max_row == 2 and wb["Пункты"].max_row > 30

    audio = client.get("/calls/1/audio")
    assert audio.status_code == 200 and audio.content == b"fake-audio"


def test_upload_text_transcript_and_manual_type(settings):
    app, client, transcriber, llm = make(settings)
    client.post(
        "/calls/upload",
        files=[("files", ("call.txt", sample_text().encode("cp1251"), "text/plain"))],
        data={"call_type": "out_request", "new_employee": "Мария Иванова"},
    )
    process_all(app)
    call = get_call(app)
    assert call.status == db.DONE and call.source == "text"
    assert call.call_type == "out_request" and call.call_type_manual and call.direction == "out"
    assert transcriber.calls == 0
    assert len(llm.requests) == 1  # тип указан вручную — без запроса на классификацию
    assert "Мария Иванова" in client.get("/calls/1").text


def test_not_target_and_short_calls(settings):
    short = Transcript(segments=[Segment(0, 1, "алло")], duration=5.0)
    app, client, _, llm = make(settings, transcriber=FakeTranscriber(short))
    client.post("/calls/upload", files=[("files", ("a.wav", b"x", "audio/wav"))])
    process_all(app)
    call = get_call(app)
    assert call.status == db.SKIPPED and "короткий" in call.status_message
    assert not llm.requests

    app2, client2, _, _ = make(settings, llm=FakeLLM(call_type="not_target"))
    client2.post("/calls/upload", files=[("files", ("b.wav", b"x", "audio/wav"))])
    process_all(app2)
    call = get_call(app2, 2)
    assert call.status == db.SKIPPED and call.status_message.startswith("Нецелевой")


def test_error_and_reevaluate(settings):
    app, client, _, llm = make(settings, llm=FakeLLM(raw=["не json", "опять не json"]))
    client.post("/calls/upload", files=[("files", ("c.txt", sample_text().encode(), "text/plain"))])
    process_all(app)
    call = get_call(app)
    assert call.status == db.ERROR and "JSON" in call.status_message
    assert "Ошибка" in client.get("/calls/1").text

    client.post("/calls/1/reevaluate")
    assert get_call(app).status == db.QUEUED
    process_all(app)
    assert get_call(app).status == db.DONE


def test_update_call_type_requeues(settings):
    app, client, _, _ = make(settings)
    client.post("/calls/upload", files=[("files", ("c.txt", sample_text().encode(), "text/plain"))])
    process_all(app)
    client.post("/calls/1/update", data={"call_type": "out_reminder", "employee_id": ""})
    call = get_call(app)
    assert call.status == db.QUEUED and call.call_type == "out_reminder" and call.call_type_manual
    process_all(app)
    assert get_call(app).status == db.DONE

    client.post("/calls/1/delete")
    assert get_call(app) is None


def test_password(settings):
    settings.app_password = "secret"
    app, client, _, _ = make(settings)
    assert client.get("/calls").status_code == 401
    assert client.get("/calls", auth=("admin", "wrong")).status_code == 401
    assert client.get("/calls", auth=("admin", "secret")).status_code == 200


def test_guess_datetime():
    assert guess_datetime("2026-10-08_14-32-05.mp3").isoformat() == "2026-10-08T14:32:05"
    assert guess_datetime("rec 20261008 1432.mp3").isoformat() == "2026-10-08T14:32:00"
    assert guess_datetime("08.10.2026 09-05.wav").isoformat() == "2026-10-08T09:05:00"
    assert guess_datetime("zapis.mp3") is None


def test_match_employee(settings):
    from app.pipeline import match_employee

    app, client, _, _ = make(settings)
    for name in ("Анна Смирнова", "Мария Иванова", "Мария Петрова"):
        client.post("/employees", data={"name": name})
    with app.state.session_factory() as s:
        clinic = app.state.clinic_id
        assert match_employee(s, clinic, "Анна").name == "Анна Смирнова"
        assert match_employee(s, clinic, "анна смирнова").name == "Анна Смирнова"
        assert match_employee(s, clinic, "Мария") is None  # две Марии — не угадываем
        assert match_employee(s, clinic, "Ольга") is None
        assert match_employee(s, clinic, None) is None
