"""Сводная статистика по администраторам и выгрузка в Excel."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session, selectinload

from . import db
from .checklists import ChecklistSet
from .scoring import NA, NO, RESULT_LABELS, YES


def period_start(days: int) -> datetime | None:
    return datetime.now() - timedelta(days=days) if days > 0 else None


def _call_time():
    return func.coalesce(db.Call.started_at, db.Call.created_at)


@dataclass
class EmployeeStats:
    employee_id: int | None
    name: str
    calls: int
    avg_total: float | None
    avg_general: float | None
    avg_scenario: float | None


@dataclass
class CriterionStats:
    criterion_id: str
    num: str
    text: str
    checklist_title: str
    failed: int
    applicable: int

    @property
    def fail_rate(self) -> float:
        return round(100.0 * self.failed / self.applicable, 1) if self.applicable else 0.0


def employee_stats(s: Session, clinic_id: int, since: datetime | None) -> list[EmployeeStats]:
    q = (
        select(
            db.Call.employee_id,
            func.count(db.Call.id),
            func.avg(db.Call.score_total),
            func.avg(db.Call.score_general),
            func.avg(db.Call.score_scenario),
        )
        .where(db.Call.clinic_id == clinic_id, db.Call.status == db.DONE)
        .group_by(db.Call.employee_id)
    )
    if since:
        q = q.where(_call_time() >= since)
    names = {e.id: e.name for e in s.scalars(select(db.Employee))}

    def r(v):
        return round(v, 1) if v is not None else None

    rows = [
        EmployeeStats(eid, names.get(eid, "Не указан"), n, r(t), r(g), r(sc)) for eid, n, t, g, sc in s.execute(q)
    ]
    return sorted(rows, key=lambda x: (x.employee_id is None, -(x.avg_total or 0)))


def criterion_stats(
    s: Session, checklists: ChecklistSet, clinic_id: int, since: datetime | None, employee_id: int | None = None
) -> list[CriterionStats]:
    q = (
        select(db.CriterionResult.criterion_id, db.CriterionResult.result, func.count())
        .join(db.Call, db.Call.id == db.CriterionResult.call_id)
        .where(db.Call.clinic_id == clinic_id, db.Call.status == db.DONE, db.CriterionResult.result != NA)
        .group_by(db.CriterionResult.criterion_id, db.CriterionResult.result)
    )
    if since:
        q = q.where(_call_time() >= since)
    if employee_id:
        q = q.where(db.Call.employee_id == employee_id)

    counts: dict[str, dict[str, int]] = {}
    for cid, result, n in s.execute(q):
        counts.setdefault(cid, {YES: 0, NO: 0})[result] = n

    titles = {c.id: checklists.general.title for c in checklists.general.criteria}
    for scenario in checklists.scenarios.values():
        titles.update({c.id: scenario.title for c in scenario.criteria})

    stats = []
    for cid, c in counts.items():
        criterion = checklists.criterion(cid)
        if criterion is None:  # пункт удалён из чек-листа
            continue
        stats.append(
            CriterionStats(cid, criterion.num, criterion.text, titles.get(cid, ""), c[NO], c[YES] + c[NO])
        )
    return sorted(stats, key=lambda x: (-x.failed, -x.fail_rate, x.criterion_id))


# ---------------------------------------------------------------- Excel

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="2F5D8A")
_RESULT_FILLS = {
    YES: PatternFill("solid", fgColor="D9F2D9"),
    NO: PatternFill("solid", fgColor="F8D7D7"),
}


def _sheet(wb: Workbook, title: str, headers: list[str], widths: list[int], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    return ws


def build_excel(s: Session, checklists: ChecklistSet, clinic_id: int, since: datetime | None) -> bytes:
    q = (
        select(db.Call)
        .where(db.Call.clinic_id == clinic_id, db.Call.status.in_((db.DONE, db.SKIPPED)))
        .options(selectinload(db.Call.results), selectinload(db.Call.employee))
        .order_by(_call_time())
    )
    if since:
        q = q.where(_call_time() >= since)
    calls = s.scalars(q).all()

    wb = Workbook()
    ws_calls = _sheet(
        wb,
        "Звонки",
        ["№ звонка", "Дата и время", "Администратор", "Тип звонка", "Длительность, с",
         "Итог, %", "Общие правила, %", "Сценарий, %", "Статус", "Файл"],
        [10, 18, 22, 40, 14, 10, 16, 12, 40, 30],
        first=True,
    )
    ws_items = _sheet(
        wb,
        "Пункты",
        ["№ звонка", "Дата и время", "Администратор", "Тип звонка", "Код", "№ пункта",
         "Пункт", "Значимость", "Оценка", "Цитата", "Комментарий"],
        [10, 18, 22, 32, 9, 9, 70, 12, 14, 50, 50],
    )

    for call in calls:
        when = call.when.strftime("%d.%m.%Y %H:%M")
        admin = call.employee.name if call.employee else "Не указан"
        ctype = checklists.scenario_title(call.call_type)
        status = "Оценён" if call.status == db.DONE else f"Не оценивается: {call.status_message}"
        ws_calls.append([
            call.id, when, admin, ctype, round(call.duration_sec or 0),
            call.score_total, call.score_general, call.score_scenario, status, call.original_filename,
        ])
        for r in call.results:
            c = checklists.criterion(r.criterion_id)
            ws_items.append([
                call.id, when, admin, ctype, r.criterion_id, c.num if c else "", c.text if c else "",
                checklists.weight_labels.get(c.weight_key, "") if c else "",
                RESULT_LABELS.get(r.result, r.result), r.evidence, r.comment,
            ])
            if r.result in _RESULT_FILLS:
                ws_items.cell(row=ws_items.max_row, column=9).fill = _RESULT_FILLS[r.result]

    ws_sum = _sheet(
        wb,
        "Сводка",
        ["Администратор", "Оценено звонков", "Итог, %", "Общие правила, %", "Сценарий, %"],
        [26, 16, 10, 16, 12],
    )
    for row in employee_stats(s, clinic_id, since):
        ws_sum.append([row.name, row.calls, row.avg_total, row.avg_general, row.avg_scenario])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
